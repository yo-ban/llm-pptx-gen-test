"""
python-pptx ライブラリを使用した PowerPoint 操作関連のヘルパー関数。
"""
import logging
from typing import List, Dict, Any, Optional
from pptx import Presentation
from pptx.text.text import TextFrame
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Pt
from models import ContentBlock

logger = logging.getLogger(__name__)


def get_placeholder_details(prs: Presentation, layout_index: int) -> List[Dict[str, Any]]:
    """
    指定されたレイアウトから利用可能なプレースホルダー情報を取得する（定型除く）。

    Args:
        prs (Presentation): Presentation オブジェクト。
        layout_index (int): レイアウトインデックス。

    Returns:
        List[Dict[str, Any]]: プレースホルダー情報のリスト。
    """
    try:
        slide_layout = prs.slide_layouts[layout_index]
    except IndexError:
        logger.error(f"レイアウトインデックス {layout_index} が範囲外です。")
        return []

    details = []
    excluded_types = [
        PP_PLACEHOLDER.DATE,
        PP_PLACEHOLDER.FOOTER,
        PP_PLACEHOLDER.SLIDE_NUMBER
    ]

    for shape in slide_layout.placeholders:
        try:
            placeholder_format = shape.placeholder_format
            if placeholder_format.type not in excluded_types:
                ph_type_str = str(placeholder_format.type)
                details.append({
                    "index": placeholder_format.idx, "name": shape.name,
                    "type": ph_type_str, "shape_id": shape.shape_id
                })
                logger.info(
                    f"  - Index: {placeholder_format.idx}, Name: '{shape.name}', Type: {ph_type_str}")
        except AttributeError:
            logger.debug(f"シェイプ '{shape.name}' はプレースホルダーフォーマットを持たないためスキップ。")
            continue
        except Exception as e:
            logger.warning(
                f"プレースホルダー情報取得エラー: Shape ID {shape.shape_id}, Error: {e}")
            continue

    details.sort(key=lambda x: x["index"])
    return details

def add_content_blocks_to_text_frame(text_frame: TextFrame, content_blocks: List[ContentBlock]):
    """
    テキストフレームにコンテンツブロックを追加する。

    Args:
        text_frame: 対象の TextFrame オブジェクト。
        content_blocks (List[ContentBlock]): 追加するコンテンツブロック情報。
    """
    text_frame.clear()
    text_frame.word_wrap = True
    first_paragraph_available = len(text_frame.paragraphs) > 0

    # フォントサイズを設定するための辞書
    font_size_by_level = {
        0: 16,
        1: 12,
        2: 10,
        3: 8
    }

    for i, content_block in enumerate(content_blocks):

        header_text = content_block.heading.strip()
        if not header_text and not content_block.items:
            logger.info("  - スキップ: 空のコンテンツブロック")
            continue

        if i == 0 and first_paragraph_available:
            p_header = text_frame.paragraphs[0]
            p_header.text = header_text
        else:
            p_header = text_frame.add_paragraph()
            p_header.text = header_text

        p_header.font.bold = True
        if content_block.font_size and content_block.font_size > 0:
            p_header.font.size = Pt(content_block.font_size)
        else:
            p_header.font.size = Pt(font_size_by_level[0])
        p_header.level = 0

        for item in content_block.items:
            item_text = item.text.strip()
            if not item_text:
                logger.info("  - スキップ: 空の箇条書きアイテム")
                continue
            p_item = text_frame.add_paragraph()
            p_item.text = item_text
            p_item.level = item.level
            if item.font_size and item.font_size > 0:
                p_item.font.size = Pt(item.font_size)
            else:
                p_item.font.size = Pt(font_size_by_level[item.level])
